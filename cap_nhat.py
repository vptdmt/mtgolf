import openpyxl, json, os

excel_file = "Up 21h00 24.10 Chia Fly giải 25.10 update title.xlsx"
if not os.path.exists(excel_file):
    excel_file = "DS khách mời giải Golf 2026- up 24.10.2026.xlsx"

wb = openpyxl.load_workbook(excel_file, data_only=True)

# 1. Đọc thông tin Size, SĐT, VIP, Phòng từ các bản ghi cũ
with open('golfers.json', 'r', encoding='utf-8') as f:
    old_data = json.load(f)
lookup = {g['name'].lower().strip(): g for g in old_data}

golfers = []
flight_groups = {}

# 2. Đọc Ca Sáng (Flight- 5h30)
if 'Flight- 5h30' in wb.sheetnames:
    ws = wb['Flight- 5h30']
    cur_fl = ''
    cur_hl = ''
    for r in range(5, 50):
        c_fl = ws.cell(r, 2).value or ws.cell(r, 3).value
        c_hl = ws.cell(r, 3).value or ws.cell(r, 4).value
        c_name = ws.cell(r, 4).value or ws.cell(r, 5).value
        c_vga = ws.cell(r, 5).value or ws.cell(r, 6).value

        # Phân tích hàng
        vals = [str(ws.cell(r, c).value or '').strip() for c in range(1, 12)]
        name = ''
        for v in vals:
            if v and v.lower() in lookup:
                name = lookup[v.lower()]['name']
                break
        
        # Lấy từ cấu trúc chuẩn
        val_name = ws.cell(r, 5).value or ws.cell(r, 4).value
        if val_name and str(val_name).strip().lower() not in ['goler name', 'họ và tên', 'tên', 'none']:
            name = str(val_name).strip()
            fl = ws.cell(r, 2).value or ws.cell(r, 3).value
            hl = ws.cell(r, 3).value or ws.cell(r, 4).value
            if fl and str(fl).strip(): cur_fl = str(fl).strip()
            if hl and str(hl).strip(): cur_hl = str(hl).strip()
            vga = str(ws.cell(r, 6).value or ws.cell(r, 5).value or '').strip()
            if vga == '0' or vga == 'None': vga = ''

            old = lookup.get(name.lower(), {})
            g = {
                'id': f'g-sang-{len(golfers)+1}',
                'stt': len(golfers)+1,
                'name': name,
                'gender': old.get('gender', ''),
                'shift': 'Sáng',
                'teeTime': '6h00 ngày 25/10/2026',
                'flight': cur_fl,
                'hole': cur_hl,
                'vga': vga or old.get('vga', ''),
                'size': old.get('size', ''),
                'title': old.get('title', ''),
                'organization': old.get('organization', ''),
                'phone': old.get('phone', ''),
                'room': old.get('room', ''),
                'vip': old.get('vip', ''),
                'pic': old.get('pic', ''),
                'note': old.get('note', ''),
                'status': old.get('status', 'Confirm'),
                'handicapIndex': old.get('handicapIndex', ''),
                'courseHandicap': old.get('courseHandicap', ''),
                'division': old.get('division', ''),
                'vhandicapNickname': old.get('vhandicapNickname', ''),
                'checkedIn': old.get('checkedIn', False),
                'checkInTime': old.get('checkInTime', None),
                'receivedShirt': old.get('receivedShirt', False),
                'receivedDoorGift': old.get('receivedDoorGift', False),
                'receivedCoupon': old.get('receivedCoupon', False),
                'teammates': []
            }
            golfers.append(g)
            k = ('Sáng', cur_fl)
            if k not in flight_groups: flight_groups[k] = []
            flight_groups[k].append(name)

# 3. Đọc Ca Chiều (Flight- 11h00)
if 'Flight- 11h00' in wb.sheetnames:
    ws = wb['Flight- 11h00']
    cur_fl = ''
    cur_hl = ''
    for r in range(4, 200):
        val_name = ws.cell(r, 6).value or ws.cell(r, 5).value
        if not val_name: continue
        name = str(val_name).strip()
        if name.lower() in ['goler name', 'họ và tên', 'tên', 'none']: continue

        fl = ws.cell(r, 2).value or ws.cell(r, 3).value
        hl = ws.cell(r, 3).value or ws.cell(r, 4).value
        if fl and str(fl).strip(): cur_fl = str(fl).strip()
        if hl and str(hl).strip(): cur_hl = str(hl).strip()

        vga = str(ws.cell(r, 7).value or ws.cell(r, 6).value or '').strip()
        if vga == '0' or vga == 'None': vga = ''
        title = str(ws.cell(r, 8).value or ws.cell(r, 7).value or '').strip()

        old = lookup.get(name.lower(), {})
        g = {
            'id': f'g-chieu-{len(golfers)+1}',
            'stt': len(golfers)+1,
            'name': name,
            'gender': old.get('gender', 'Ông'),
            'shift': 'Chiều',
            'teeTime': '11h30 ngày 25/10/2026',
            'flight': cur_fl,
            'hole': cur_hl,
            'vga': vga or old.get('vga', ''),
            'size': old.get('size', ''),
            'title': title or old.get('title', ''),
            'organization': old.get('organization', ''),
            'phone': old.get('phone', ''),
            'room': old.get('room', ''),
            'vip': old.get('vip', ''),
            'pic': old.get('pic', ''),
            'note': old.get('note', ''),
            'status': old.get('status', 'Confirm'),
            'handicapIndex': old.get('handicapIndex', ''),
            'courseHandicap': old.get('courseHandicap', ''),
            'division': old.get('division', ''),
            'vhandicapNickname': old.get('vhandicapNickname', ''),
            'checkedIn': old.get('checkedIn', False),
            'checkInTime': old.get('checkInTime', None),
            'receivedShirt': old.get('receivedShirt', False),
            'receivedDoorGift': old.get('receivedDoorGift', False),
            'receivedCoupon': old.get('receivedCoupon', False),
            'teammates': []
        }
        golfers.append(g)
        k = ('Chiều', cur_fl)
        if k not in flight_groups: flight_groups[k] = []
        flight_groups[k].append(name)

# 4. Cập nhật bạn cùng Flight
for g in golfers:
    k = (g['shift'], g['flight'])
    if k in flight_groups:
        g['teammates'] = [m for m in flight_groups[k] if m != g['name']]

# Lưu ra JSON và data.js
with open('golfers.json', 'w', encoding='utf-8') as f:
    json.dump(golfers, f, ensure_ascii=False, indent=2)

with open('data.js', 'w', encoding='utf-8') as f:
    f.write('window.INITIAL_GOLFERS = ' + json.dumps(golfers, ensure_ascii=False, indent=2) + ';\n')

print(f'Thanh cong! Da cap nhat {len(golfers)} golfer.')
