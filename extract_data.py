import openpyxl, json, re, sys

sys.stdout.reconfigure(encoding='utf-8')
excel_path = 'DS khách mời giải Golf 2026- up 24.10.2026.xlsx'
wb = openpyxl.load_workbook(excel_path, data_only=True)

# 1. Parse DS sheet
ds_sheet = wb['DS']
ds_lookup = {}
for r in range(8, 250):
    val_name = ds_sheet.cell(r, 5).value
    if not val_name:
        continue
    name_clean = str(val_name).strip()
    key = name_clean.lower()
    
    # Clean phone
    phone_raw = str(ds_sheet.cell(r, 13).value or '').strip()
    if phone_raw.endswith('.0'):
        phone_raw = phone_raw[:-2]
    if phone_raw and not phone_raw.startswith('0') and phone_raw.isdigit():
        phone_raw = '0' + phone_raw
        
    vga_raw = str(ds_sheet.cell(r, 6).value or '').strip()
    if vga_raw.endswith('.0'):
        vga_raw = vga_raw[:-2]

    ds_lookup[key] = {
        'gender': str(ds_sheet.cell(r, 4).value or '').strip(),
        'name': name_clean,
        'vga': vga_raw,
        'title': str(ds_sheet.cell(r, 11).value or '').strip(),
        'organization': str(ds_sheet.cell(r, 12).value or '').strip(),
        'phone': phone_raw,
        'size': str(ds_sheet.cell(r, 14).value or '').strip(),
        'room': str(ds_sheet.cell(r, 16).value or '').strip(),
        'vip': str(ds_sheet.cell(r, 17).value or '').strip(),
        'note': str(ds_sheet.cell(r, 18).value or '').strip(),
        'pic': str(ds_sheet.cell(r, 20).value or '').strip(),
        'status': str(ds_sheet.cell(r, 21).value or '').strip(),
        'handicapIndex': str(ds_sheet.cell(r, 7).value or '').strip(),
        'courseHandicap': str(ds_sheet.cell(r, 8).value or '').strip(),
        'vhandicapNickname': str(ds_sheet.cell(r, 10).value or '').strip()
    }

golfers = []
flight_groups = {} # (shift, flight) -> list of golfer names

# Helper to clean VGA & numbers
def clean_num(val):
    if val is None:
        return ''
    s = str(val).strip()
    if s.endswith('.0'):
        s = s[:-2]
    return s

# 2. Parse Flight- 5h30 (Ca Sáng)
ws1 = wb['Flight- 5h30']
cur_flight = ''
cur_hole = ''
for r in range(6, 60):
    val_name = ws1.cell(r, 5).value
    val_fl = ws1.cell(r, 3).value
    val_hl = ws1.cell(r, 4).value
    if val_fl is not None and str(val_fl).strip() != '':
        cur_flight = clean_num(val_fl)
    if val_hl is not None and str(val_hl).strip() != '':
        cur_hole = str(val_hl).strip()
    if not val_name:
        continue
    name = str(val_name).strip()
    if name.lower() in ['goler name', 'họ và tên', 'họ tên', 'tên', 'none']:
        continue
        
    ds_info = ds_lookup.get(name.lower(), {})
    gender = str(ws1.cell(r, 2).value or ds_info.get('gender') or '').strip()
    if gender in ['Sáng', 'Chiều']:
        gender = ds_info.get('gender', '')
    
    vga = clean_num(ws1.cell(r, 6).value) or ds_info.get('vga', '')
    size = str(ws1.cell(r, 7).value or ds_info.get('size', '') or '').strip()
    hdc = clean_num(ws1.cell(r, 8).value) or ds_info.get('handicapIndex', '')
    chdc = clean_num(ws1.cell(r, 9).value) or ds_info.get('courseHandicap', '')
    div = str(ws1.cell(r, 10).value or '').strip()
    nick = str(ws1.cell(r, 11).value or ds_info.get('vhandicapNickname', '') or '').strip()
    
    g_id = f"g-sang-{r}"
    g_data = {
        'id': g_id,
        'stt': len(golfers) + 1,
        'name': name,
        'gender': gender,
        'shift': 'Sáng',
        'teeTime': '6h00 ngày 25/10/2026',
        'flight': cur_flight,
        'hole': cur_hole,
        'vga': vga,
        'size': size,
        'title': ds_info.get('title', ''),
        'organization': ds_info.get('organization', ''),
        'phone': ds_info.get('phone', ''),
        'room': ds_info.get('room', ''),
        'vip': ds_info.get('vip', ''),
        'pic': ds_info.get('pic', ''),
        'note': ds_info.get('note', ''),
        'status': ds_info.get('status', 'Confirm'),
        'handicapIndex': hdc,
        'courseHandicap': chdc,
        'division': div,
        'vhandicapNickname': nick,
        'checkedIn': False,
        'checkInTime': None,
        'receivedShirt': False,
        'receivedDoorGift': False,
        'receivedCoupon': False,
        'teammates': []
    }
    golfers.append(g_data)
    key = ('Sáng', cur_flight)
    if key not in flight_groups:
        flight_groups[key] = []
    flight_groups[key].append(name)

# 3. Parse Flight- 11h00 (Ca Chiều)
ws2 = wb['Flight- 11h00']
cur_flight = ''
cur_hole = ''
for r in range(5, 200):
    val_name = ws2.cell(r, 6).value
    val_fl = ws2.cell(r, 3).value
    val_hl = ws2.cell(r, 4).value
    if val_fl is not None and str(val_fl).strip() != '':
        cur_flight = clean_num(val_fl)
    if val_hl is not None and str(val_hl).strip() != '':
        cur_hole = str(val_hl).strip()
    if not val_name:
        continue
    name = str(val_name).strip()
    if name.lower() in ['goler name', 'họ và tên', 'họ tên', 'tên', 'none']:
        continue
        
    ds_info = ds_lookup.get(name.lower(), {})
    gender = str(ws2.cell(r, 5).value or ds_info.get('gender') or '').strip()
    vga = clean_num(ws2.cell(r, 7).value) or ds_info.get('vga', '')
    size = str(ws2.cell(r, 8).value or ds_info.get('size', '') or '').strip()
    title = str(ws2.cell(r, 9).value or ds_info.get('title', '') or '').strip()
    hdc = clean_num(ws2.cell(r, 10).value) or ds_info.get('handicapIndex', '')
    chdc = clean_num(ws2.cell(r, 11).value) or ds_info.get('courseHandicap', '')
    div = str(ws2.cell(r, 12).value or '').strip()
    nick = str(ws2.cell(r, 13).value or ds_info.get('vhandicapNickname', '') or '').strip()
    note_flight = str(ws2.cell(r, 14).value or '').strip()
    full_note = ds_info.get('note', '')
    if note_flight and note_flight not in full_note:
        full_note = (full_note + ' | ' + note_flight).strip(' | ')

    g_id = f"g-chieu-{r}"
    g_data = {
        'id': g_id,
        'stt': len(golfers) + 1,
        'name': name,
        'gender': gender,
        'shift': 'Chiều',
        'teeTime': '11h30 ngày 25/10/2026',
        'flight': cur_flight,
        'hole': cur_hole,
        'vga': vga,
        'size': size,
        'title': title,
        'organization': ds_info.get('organization', ''),
        'phone': ds_info.get('phone', ''),
        'room': ds_info.get('room', ''),
        'vip': ds_info.get('vip', ''),
        'pic': ds_info.get('pic', ''),
        'note': full_note,
        'status': ds_info.get('status', 'Confirm'),
        'handicapIndex': hdc,
        'courseHandicap': chdc,
        'division': div,
        'vhandicapNickname': nick,
        'checkedIn': False,
        'checkInTime': None,
        'receivedShirt': False,
        'receivedDoorGift': False,
        'receivedCoupon': False,
        'teammates': []
    }
    golfers.append(g_data)
    key = ('Chiều', cur_flight)
    if key not in flight_groups:
        flight_groups[key] = []
    flight_groups[key].append(name)

# 4. Check for guests in DS not in flights
names_in_flights = {g['name'].lower() for g in golfers}
for key, ds_info in ds_lookup.items():
    if key not in names_in_flights and ds_info['name'] and ds_info['status'] != 'Hủy':
        g_id = f"g-ds-{len(golfers)+1}"
        golfers.append({
            'id': g_id,
            'stt': len(golfers) + 1,
            'name': ds_info['name'],
            'gender': ds_info['gender'],
            'shift': 'Chưa xếp',
            'teeTime': 'Chưa xếp giờ',
            'flight': '',
            'hole': '',
            'vga': ds_info['vga'],
            'size': ds_info['size'],
            'title': ds_info['title'],
            'organization': ds_info['organization'],
            'phone': ds_info['phone'],
            'room': ds_info['room'],
            'vip': ds_info['vip'],
            'pic': ds_info['pic'],
            'note': ds_info['note'],
            'status': ds_info['status'] or 'Dự bị/Chưa xếp',
            'handicapIndex': ds_info['handicapIndex'],
            'courseHandicap': ds_info['courseHandicap'],
            'division': '',
            'vhandicapNickname': ds_info['vhandicapNickname'],
            'checkedIn': False,
            'checkInTime': None,
            'receivedShirt': False,
            'receivedDoorGift': False,
            'receivedCoupon': False,
            'teammates': []
        })

# 5. Populate teammates
for g in golfers:
    shift = g['shift']
    fl = g['flight']
    if fl and (shift, fl) in flight_groups:
        all_mates = flight_groups[(shift, fl)]
        g['teammates'] = [m for m in all_mates if m != g['name']]

with open('golfers.json', 'w', encoding='utf-8') as f:
    json.dump(golfers, f, ensure_ascii=False, indent=2)

print(f'Successfully exported {len(golfers)} golfers to golfers.json')
sang_count = sum(1 for g in golfers if g['shift'] == 'Sáng')
chieu_count = sum(1 for g in golfers if g['shift'] == 'Chiều')
other_count = sum(1 for g in golfers if g['shift'] == 'Chưa xếp')
print(f'Ca Sáng: {sang_count}, Ca Chiều: {chieu_count}, Chưa xếp: {other_count}')
